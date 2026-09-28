"""Read-only workbook extraction. Never writes back to the source workbook.

Run with the Codex bundled Python runtime (openpyxl). Outputs contain customer
data and belong in ignored private-data/, never public assets or source control.
"""
import argparse
import hashlib
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
import openpyxl


def value(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    if hasattr(v, "text"):
        return {"formula": v.text, "range": v.ref}
    return v


def inspect(source, output):
    formulas = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True)
    fingerprint = hashlib.sha256(Path(source).read_bytes()).hexdigest()
    sheets, issues, source_rows = [], [], {}
    for sheet in formulas:
        rows = []
        counts = Counter()
        for row in sheet:
            cells = {}
            for cell in row:
                if cell.value is None:
                    continue
                datum = {"value": value(cell.value), "cached": value(cached[sheet.title][cell.coordinate].value)}
                if cell.hyperlink:
                    datum["link"] = cell.hyperlink.target
                    counts["links"] += 1
                if cell.comment:
                    datum["comment"] = cell.comment.text
                if cell.data_type == "f" or hasattr(cell.value, "text"):
                    counts["formulas"] += 1
                if cached[sheet.title][cell.coordinate].data_type == "e":
                    counts["cached_errors"] += 1
                if isinstance(cell.value, str) and "#REF!" in cell.value:
                    counts["broken_reference_formulas"] += 1
                cells[cell.coordinate] = datum
            if cells:
                rows.append({"row": row[0].row, "cells": cells})
        source_rows[sheet.title] = rows
        sheets.append({"name": sheet.title, "rows": len(rows), "columns": sheet.max_column, **counts})

    crm = []
    for row in cached["GetDataFromRecruitCRM"].iter_rows(min_row=2):
        vals = [c.value for c in row]
        if vals[0] is not None:
            crm.append({"external_id": str(vals[0]), "client": vals[1], "title": vals[2], "created_date": value(vals[3]), "status": vals[4], "source_row": row[0].row})
    crm_ids = {r["external_id"] for r in crm}
    tracker = cached["Sourcing Tracker"]
    sessions = []
    staff = set()
    for r in range(3, tracker.max_row + 1):
        get = lambda col: tracker[f"{col}{r}"].value
        if not get("A") or not get("E"):
            continue
        if not isinstance(get("A"), datetime):
            issues.append({"kind": "invalid_work_date", "sheet": tracker.title, "row": r})
            continue
        label = str(get("E"))
        external_id = label.split("|")[0].strip() if "|" in label else None
        if external_id not in crm_ids:
            issues.append({"kind": "unmatched_search", "sheet": tracker.title, "row": r, "search_label": label})
        entries = []
        for name_col, mapped_col, peer_col, partner_col in [("O", "P", "Q", "R"), ("T", "U", "V", "W"), ("Y", "Z", "AA", "AB"), ("AC", "AD", "AE", "AF")]:
            name = get(name_col)
            if not name and all(get(c) is None for c in [mapped_col, peer_col, partner_col]):
                continue
            entry = {"staff_name": name, "mapped": get(mapped_col), "peer_approved": get(peer_col), "partner_approved": get(partner_col), "source_cells": [f"{c}{r}" for c in [name_col, mapped_col, peer_col, partner_col]]}
            if name:
                staff.add(str(name).strip())
            if all(isinstance(entry[k], (int, float)) for k in ["mapped", "peer_approved", "partner_approved"]):
                if not 0 <= entry["partner_approved"] <= entry["peer_approved"] <= entry["mapped"]:
                    issues.append({"kind": "non_monotonic_counts", "sheet": tracker.title, "row": r, "cells": entry["source_cells"]})
            entries.append(entry)
        links = []
        for col in ["AI", "E"]:
            c = formulas[tracker.title][f"{col}{r}"]
            if c.hyperlink and c.hyperlink.target not in links:
                links.append(c.hyperlink.target)
        totals = {"mapped": get("I"), "peer_approved": get("J"), "partner_approved": get("M")}
        for metric, reported in totals.items():
            known = [e[metric] for e in entries if isinstance(e[metric], (int, float))]
            if known and isinstance(reported, (int, float)) and sum(known) != reported:
                issues.append({"kind": "aggregate_mismatch", "sheet": tracker.title, "row": r, "metric": metric, "cached_total": reported, "individual_sum": sum(known)})
        sessions.append({"source_row": r, "date": value(get("A")), "team": get("C"), "client": get("D"), "search_label": label, "external_id": external_id, "session_number": get("F"), "target_approved": get("G"), "readiness": get("H"), "partner_label": get("K"), "review_date": value(get("L")), "handover_date": value(get("N")), "source_totals": totals, "unlabeled_AG_value": get("AG"), "breakdown_links": links, "entries": entries})
    by_slot = Counter((s["date"], s["team"]) for s in sessions)
    duplicate_slots = [{"date": d, "team": t, "sessions": n} for (d, t), n in by_slot.items() if n > 1]
    out = {"source_sha256": fingerprint, "source_filename": Path(source).name, "sheets": sheets, "crm_searches": crm, "sessions": sessions, "staff_names": sorted(staff), "issues": issues, "duplicate_team_days": duplicate_slots, "raw_rows": source_rows}
    output.mkdir(parents=True, exist_ok=True)
    (output / "workbook-extract.json").write_text(json.dumps(out, indent=2, default=str))
    summary = {"source_sha256": fingerprint, "sheets": sheets, "crm_searches": len(crm), "sessions": len(sessions), "staff_names": len(staff), "issue_counts": dict(Counter(i["kind"] for i in issues)), "duplicate_team_days": len(duplicate_slots), "date_range": [min(s["date"] for s in sessions), max(s["date"] for s in sessions)]}
    (output / "inspection-summary.json").write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("private-data"))
    args = parser.parse_args()
    inspect(args.source, args.output)

