from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
wb=Workbook();guide=wb.active;guide.title='Instructions'
rows=[['XQtiv Search and Candidate Import'],['1. Fill Searches: Search name and Client are required. Status defaults to Open.'],['2. Supply a unique XQtiv Search ID for every search with associated candidates.'],['3. Fill Candidates: repeat that same XQtiv Search ID on each associated candidate row.'],['4. Required candidate fields: First Name, Last Name, LinkedIn URL. XQtiv Search ID is required to link a search.'],['5. The same LinkedIn profile can appear on different searches; it is reused, not duplicated.'],['6. Upload in Integrations > All searches > Import Excel, or Candidates > Import Excel. Preview, then confirm.'],['7. Limits: 200 searches and 500 candidate rows per workbook.'],['8. Searches with blank IDs receive automatic IDs, but cannot be linked from Candidates in this upload.'],['9. Existing search matches and existing candidate mappings are skipped; existing data is not overwritten.'],['10. Imported searches are managed in this app. New mappings start as drafts under the importing researcher.'],['11. This template does not import approval decisions, engagement stages, dated activity history or attachments.'],['12. Candidate import requires Super admin or Data quality analyst plus Researcher permissions; search import requires Admin.'],['For candidate-only imports, use the Candidates tab and leave Searches empty. Search IDs are optional there.'],['Example Searches row: 201 | Head of Sales | Example Client | Closed'],['Example Candidates row: 201 | Alex | Example | https://www.linkedin.com/in/alex-example | ...'],['Examples are illustrative only; enter your actual data on the two data sheets.']]
for r in rows:guide.append(r)
guide.column_dimensions['A'].width=125
for row in guide:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
for i in range(1,len(rows)+1):guide.row_dimensions[i].height=32
s=wb.create_sheet('Searches');s.append(['XQtiv Search ID','Search name','Client','Status'])
c=wb.create_sheet('Candidates');c.append(['XQtiv Search ID','First Name','Last Name','LinkedIn URL','Email','Phone','Current Title','Current Company','Rationale'])
for ws in [s,c]:
 ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{ws.cell(1,ws.max_column).column_letter}1'
 for col in ws[1]:ws.column_dimensions[col.column_letter].width=26 if col.column<4 else 42
 for row in ws.iter_rows(min_row=2,max_row=201 if ws==s else 501):
  for cell in row:cell.number_format='@'
for ws in wb:
 for cell in ws[1]:cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=12);cell.fill=PatternFill('solid',fgColor='001B50')
 ws.sheet_view.showGridLines=True
validation=DataValidation(type='list',formula1='"Open,On Hold,Closed,Abandoned,Canceled"',allow_blank=True);validation.errorTitle='Search status';validation.error='Choose a listed search status.';validation.showErrorMessage=True;s.add_data_validation(validation);validation.add('D2:D201')
p=Path(__file__).resolve().parents[1]/'public/templates/xqtiv-search-candidate-import.xlsx';p.parent.mkdir(exist_ok=True);wb.save(p)
